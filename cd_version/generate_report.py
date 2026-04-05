import os
from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

def create_report():
    doc = Document()
    
    # Title
    title = doc.add_heading('Multi-Robot Warehouse Navigation & Task Allocation Report', 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    # Section 1: Theory
    doc.add_heading('1. Theory', level=1)
    doc.add_paragraph(
        "Multi-robot systems in warehouse environments involve managing multiple autonomous "
        "agents that must navigate shared spaces to complete tasks efficiently without colliding. "
        "The problem is generally divided into task allocation (assigning tasks to robots) and "
        "path planning (finding collision-free paths for robots to execute their tasks)."
    )
    
    # Section 2: Algorithm Details
    doc.add_heading('2. Algorithm Details', level=1)
    
    doc.add_heading('2.1 Space-Time A*', level=2)
    doc.add_paragraph(
        "Space-Time A* is an extension of the classic A* search algorithm that includes time "
        "as an additional dimension in the search space. This allows the algorithm to find paths "
        "that avoid dynamic obstacles or other robots whose future positions are known. "
        "Each node in the search tree represents a (x, y, t) state."
    )
    
    doc.add_heading('2.2 Prioritised Planning', level=2)
    doc.add_paragraph(
        "Prioritised Planning is a decoupled approach to multi-agent pathfinding. Robots are "
        "assigned unique priorities. Paths are planned one by one in decreasing order of priority. "
        "When planning for a robot, the paths of all higher-priority robots are treated as dynamic "
        "obstacles in space-time. This method is much faster than coupled approaches, although it "
        "does not guarantee completeness or optimality."
    )
    
    doc.add_heading('2.3 Task Allocation', level=2)
    doc.add_paragraph(
        "Task allocation is handled using the Hungarian Algorithm (or similar assignment methods) "
        "to continuously pair available robots with unassigned tasks based on a cost matrix. "
        "The cost is typically the heuristic distance or estimated driving time from the robot's "
        "current position to the task location."
    )
    
    doc.add_page_break()
    
    # Section 3: Mechanical Engineering Integration
    doc.add_heading('3. Mechanical Engineering Integration', level=1)
    doc.add_paragraph(
        "The algorithms implemented must account for kinematic constraints inherent in mechanical "
        "systems. This includes limiting maximum velocities, considering acceleration/deceleration "
        "profiles, and accommodating turning radii. The simulation strictly enforces orthogonal "
        "movements (no diagonal movement) to simulate differential drive or holonomic robots operating "
        "in a grid-based warehouse."
    )
    
    # Section 4: Bugs Fixed
    doc.add_heading('4. Bugs Fixed', level=1)
    doc.add_paragraph(
        "Several critical issues were addressed during development:\n"
        "- Path Collisions: Fixed issues where Space-Time A* would occasionally allow "
        "  vertex or edge collisions if time steps were not properly synchronized.\n"
        "- Priority Deadlocks: Addressed edge cases in Prioritised Planning where lower-priority "
        "  robots could get permanently blocked leading to failed path searches.\n"
        "- Task Reassignment: Fixed a bug where completed tasks were occasionally being reassigned."
    )
    
    doc.add_page_break()
    
    # Section 5: Exact Execution Output
    doc.add_heading('5. Simulation Execution Output', level=1)
    doc.add_paragraph(
        "The latest simulation run produced the following metrics:\n\n"
        "Simulation completed in 61 steps.\n\n"
        "--- Final Performance Metrics ---\n"
        "Total execution steps: 61\n"
        "Total distance traveled: 192\n"
        "Tasks completed: 9\n"
        "Average makespan (steps per task): 29.22\n"
        "Task completion rate (throughput): 14.75 tasks/100 steps\n"
        "---------------------------------\n"
        "Robot 0: Status=IDLE, Current Task=None, Valid Path=True\n"
        "Robot 1: Status=IDLE, Current Task=None, Valid Path=True\n"
        "Robot 2: Status=IDLE, Current Task=None, Valid Path=True\n"
        "Robot 3: Status=IDLE, Current Task=None, Valid Path=True\n"
    )

    # Save document
    output_path = os.path.join(os.path.dirname(__file__), 'Warehouse_Simulation_Report.docx')
    doc.save(output_path)
    print(f"Report successfully saved to {output_path}")

if __name__ == '__main__':
    create_report()